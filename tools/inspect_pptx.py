import os
from pptx import Presentation

path = r"C:\Users\HP VICTUS\Downloads\ML Review 2-Deepfake Detector.pptx"
prs = Presentation(path)

with open("pptx_dump.txt", "w", encoding="utf-8") as out:
    for idx, slide in enumerate(prs.slides):
        out.write(f"\n==================== SLIDE {idx + 1} ====================\n")
        out.write(f"Slide layout: {slide.slide_layout.name}\n")
        for shape in slide.shapes:
            out.write(f"Shape: {shape.name} (type={shape.shape_type}, id={shape.shape_id}) left={shape.left} top={shape.top} w={shape.width} h={shape.height}\n")
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    p_text = "".join([r.text for r in p.runs]) if p.runs else p.text
                    out.write(f"   P{p_idx} (level={p.level}): {p_text}\n")
                    for r_idx, r in enumerate(p.runs):
                        fn = r.font.name
                        fs = r.font.size.pt if r.font.size else None
                        fc = None
                        try:
                            if r.font.color:
                                if r.font.color.type == 1: # RGB
                                    fc = str(r.font.color.rgb)
                                elif r.font.color.type == 2: # Theme
                                    fc = f"theme:{r.font.color.theme_color}"
                        except Exception as e:
                            fc = str(e)
                        out.write(f"      Run {r_idx}: '{r.text}' | font={fn}, size={fs}, color={fc}, bold={r.font.bold}, italic={r.font.italic}\n")
            elif shape.has_table:
                table = shape.table
                out.write(f"   TABLE {len(table.rows)}x{len(table.columns)}:\n")
                for r_idx, row in enumerate(table.rows):
                    row_txt = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(f"     Row {r_idx}: {row_txt}\n")

print("Dumped to pptx_dump.txt successfully")
