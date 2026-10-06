import os, sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

path = r"C:\Users\HP VICTUS\Downloads\ML_Review_2_Missed_Lead_Detector.pptx"
print("File exists:", os.path.exists(path))
print("File size:", os.path.getsize(path), "bytes")

prs = Presentation(path)
print("Slides:", len(prs.slides))

for i, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {i+1} ===")
    shape_count = 0
    pic_count = 0
    for s in slide.shapes:
        shape_count += 1
        if s.shape_type == 13:
            pic_count += 1
            print(f"  [Picture] {s.name} pos=({s.left},{s.top}) size=({s.width},{s.height})")
        elif s.has_text_frame:
            txt = s.text_frame.text.replace('\n', ' | ')[:100]
            if txt.strip():
                print(f"  [{s.name}] {txt}")
    print(f"  Total shapes: {shape_count}, Pictures: {pic_count}")

# Check master/background watermark still exists
master = prs.slide_masters[0]
print(f"\nMaster shapes: {len(master.shapes)}")
for s in master.shapes:
    print(f"  Master shape: {s.name} type={s.shape_type}")
