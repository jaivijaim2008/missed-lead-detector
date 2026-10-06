import docx

doc = docx.Document('template_report.docx')

def update_cell_text(cell, new_text):
    if not cell.paragraphs:
        cell.text = new_text
        return
    p = cell.paragraphs[0]
    if p.runs:
        r0 = p.runs[0]
        r0.text = new_text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.text = new_text
    for extra_p in cell.paragraphs[1:]:
        extra_p.text = ""

table = doc.tables[4] # Symbols & Abbreviations table
print("Before:", table.cell(1, 0).text, "|", table.cell(1, 1).text)
update_cell_text(table.cell(1, 0), "NLP")
update_cell_text(table.cell(1, 1), "Natural Language Processing")
print("After:", table.cell(1, 0).text, "|", table.cell(1, 1).text)
